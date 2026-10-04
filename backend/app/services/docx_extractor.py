import os
import logging
from typing import List, Dict, Any
import docx

logger = logging.getLogger(__name__)

def extract_text_from_docx(docx_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a DOCX file using python-docx.
    Divides content into logical pages/sections based on page breaks or length.
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"DOCX file not found at {docx_path}")

    try:
        doc = docx.Document(docx_path)
        
        pages_data: List[Dict[str, Any]] = []
        current_page_text: List[str] = []
        current_page = 1
        current_char_count = 0
        PAGE_CHAR_THRESHOLD = 1500  # approximate standard page length

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if this paragraph contains a page break
            has_page_break = any(
                'w:br' in run._r.xml and 'w:type="page"' in run._r.xml
                for run in para.runs
            )

            current_page_text.append(text)
            current_char_count += len(text)

            if has_page_break or current_char_count >= PAGE_CHAR_THRESHOLD:
                pages_data.append({
                    "page_number": current_page,
                    "text": "\n\n".join(current_page_text)
                })
                current_page += 1
                current_page_text = []
                current_char_count = 0

        # Also extract table cells
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)
            if table_text:
                current_page_text.append("\n".join(table_text))

        if current_page_text:
            pages_data.append({
                "page_number": current_page,
                "text": "\n\n".join(current_page_text)
            })

        if not pages_data:
            pages_data.append({
                "page_number": 1,
                "text": "[Empty DOCX document]"
            })

        return pages_data

    except Exception as e:
        logger.error(f"Error reading DOCX {docx_path}: {e}")
        raise ValueError(f"Corrupted or invalid DOCX file: {str(e)}")
