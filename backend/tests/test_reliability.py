import json
import os
import subprocess
import sys
import tempfile
import unittest
import asyncio
import chromadb
from pathlib import Path
from unittest.mock import patch

import docx
import pymupdf
from fastapi import HTTPException
from PIL import Image, ImageDraw, ImageFont

from app.config import settings
from app.models.document import DocumentChunk
from app.models.investigation import SourceReference
from app.routes import system as system_routes
from app.services.answer_generator import AnswerGenerator, AnswerProviderError
from app.services import docx_extractor, ocr, pdf_extractor
from app.services.chunker import chunk_document_pages
from app.services.conflict_detector import ConflictDetector
from app.services.embeddings import (
    EmbeddingConfigurationError,
    SemanticDenseEmbeddingFunction,
    embedding_function,
)
from app.services.retriever import Retriever
from app.services.vector_store import VectorStore


ROOT = Path(__file__).resolve().parents[2]
PYTHON = sys.executable


def reference(name: str, text: str, chunk_id: str) -> SourceReference:
    return SourceReference(
        document_id=name,
        filename=name,
        page_number=1,
        chunk_id=chunk_id,
        text=text,
        file_type="txt",
        relevance_score=0.9,
    )


class ExtractionAndIndexTests(unittest.TestCase):
    def test_existing_l2_collection_is_adopted_without_changing_its_metric(self):
        with tempfile.TemporaryDirectory() as directory:
            persist_dir = str(Path(directory) / "chroma")
            client = chromadb.PersistentClient(
                path=persist_dir,
                settings=chromadb.config.Settings(anonymized_telemetry=False),
            )
            text = "Annual paid leave allowance is 20 days."
            collection = client.create_collection(
                name="legacy_l2_test",
                metadata={"hnsw:space": "l2"},
            )
            collection.add(
                ids=["legacy-chunk"],
                documents=[text],
                metadatas=[{
                    "document_id": "legacy-doc",
                    "filename": "legacy.txt",
                    "page_number": 1,
                    "file_type": "txt",
                }],
                embeddings=embedding_function([text]),
            )

            store = VectorStore(
                persist_dir=persist_dir,
                collection_name="legacy_l2_test",
            )
            results = store.query(text, n_results=1)
            raw_result = store.collection.query(
                query_texts=[text],
                n_results=1,
                include=["distances"],
            )
            distance = raw_result["distances"][0][0]

        self.assertIsNone(store.configuration_error)
        self.assertEqual(store.distance_metric, "l2")
        self.assertEqual(store.collection.metadata.get("hnsw:space"), "l2")
        self.assertAlmostEqual(
            results[0]["similarity_score"],
            max(0.0, min(1.0, 1.0 - distance / 2.0)),
            places=4,
        )

    def test_pdf_docx_and_txt_extraction_and_chunking(self):
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = Path(directory) / "policy.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "Annual leave allowance is 20 days.")
            pdf.save(pdf_path)
            pdf.close()

            docx_path = Path(directory) / "handbook.docx"
            word_doc = docx.Document()
            word_doc.add_paragraph("Annual leave allowance is 18 days.")
            word_doc.save(docx_path)

            txt_path = Path(directory) / "notice.txt"
            txt_path.write_text("Annual leave allowance is 20 days.", encoding="utf-8")

            pdf_pages = pdf_extractor.extract_text_from_pdf(str(pdf_path))
            docx_pages = docx_extractor.extract_text_from_docx(str(docx_path))
            txt_text = txt_path.read_text(encoding="utf-8")

        self.assertIn("20 days", pdf_pages[0]["text"])
        self.assertIn("18 days", docx_pages[0]["text"])
        self.assertIn("20 days", txt_text)
        self.assertEqual(docx_pages[0]["page_number"], 1)

        chunks = chunk_document_pages(
            "test-doc", "notice.txt", "txt",
            [{"page_number": 1, "text": "First sentence. Second sentence."}],
            chunk_size=20,
            chunk_overlap=0,
        )
        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(chunk.document_id == "test-doc" for chunk in chunks))
        long_chunks = chunk_document_pages(
            "long-doc",
            "large.txt",
            "txt",
            [{"page_number": 1, "text": " ".join(["substantive"] * 40)}],
            chunk_size=50,
            chunk_overlap=10,
        )
        self.assertGreater(len(long_chunks), 1)
        self.assertTrue(all(0 < len(chunk.text) <= 50 for chunk in long_chunks))

    def test_embedding_index_and_retrieval_preserve_source_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStore(
                persist_dir=str(Path(directory) / "chroma"),
                collection_name="reliability_test",
            )
            chunk = DocumentChunk(
                chunk_id="source-chunk-1",
                document_id="source-doc",
                filename="policy.txt",
                page_number=3,
                text="Annual paid leave allowance is 20 days.",
                file_type="txt",
                metadata={},
            )
            store.add_chunks([chunk])
            results = store.query("annual paid leave", n_results=1)

        self.assertEqual(results[0]["chunk_id"], "source-chunk-1")
        self.assertEqual(results[0]["filename"], "policy.txt")
        self.assertEqual(results[0]["page_number"], 3)
        self.assertGreaterEqual(results[0]["similarity_score"], 0.0)
        self.assertEqual(len(embedding_function(["test"])[0]), 384)
        self.assertEqual(
            SemanticDenseEmbeddingFunction.build_from_config(
                embedding_function.get_config()
            ).profile,
            embedding_function.profile,
        )

    def test_retrieval_supports_all_and_single_document_search_and_abstains(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStore(
                persist_dir=str(Path(directory) / "chroma"),
                collection_name="retrieval_scope_test",
            )
            store.add_chunks([
                DocumentChunk(
                    chunk_id="leave-a",
                    document_id="policy-a",
                    filename="Company_Policy.txt",
                    page_number=1,
                    text="Annual paid leave allowance is 20 days.",
                    file_type="txt",
                    metadata={},
                ),
                DocumentChunk(
                    chunk_id="leave-b",
                    document_id="policy-b",
                    filename="Employee_Handbook.txt",
                    page_number=1,
                    text="Annual paid leave allowance is 18 days.",
                    file_type="txt",
                    metadata={},
                ),
                DocumentChunk(
                    chunk_id="stipend",
                    document_id="remote",
                    filename="Remote_Work.txt",
                    page_number=1,
                    text="Remote employees receive a home office equipment stipend.",
                    file_type="txt",
                    metadata={},
                ),
            ])
            retriever = Retriever()
            retriever.vector_store = store
            exact = retriever.retrieve("What is the annual paid leave allowance?")
            judge_prompt = retriever.retrieve(
                "How many paid leaves can an employee take per year?"
            )
            paraphrase = retriever.retrieve("How many annual paid days off are granted?")
            scoped = retriever.retrieve(
                "What is the annual paid leave allowance?",
                document_ids=["policy-a"],
            )
            irrelevant = retriever.retrieve("What is the pet insurance policy?")

        self.assertEqual(
            {source.document_id for source in exact},
            {"policy-a", "policy-b"},
        )
        self.assertEqual(
            {source.document_id for source in judge_prompt},
            {"policy-a", "policy-b"},
        )
        conflict_answer = AnswerGenerator().generate_investigation(
            "judge-prompt",
            "How many paid leaves can an employee take per year?",
            judge_prompt,
            "now",
        )
        self.assertTrue(conflict_answer.conflict_detected)
        self.assertEqual(conflict_answer.confidence, "LOW")
        self.assertEqual(len(conflict_answer.sources), 2)
        self.assertTrue(paraphrase)
        self.assertEqual({source.document_id for source in scoped}, {"policy-a"})
        self.assertEqual(irrelevant, [])

    def test_ocr_success_is_real_text_and_unavailable_is_explicit(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "notice.jpg"
            Image.new("RGB", (10, 10), "white").save(image_path)
            with patch.object(ocr, "ensure_tesseract_available"), patch.object(
                ocr.pytesseract, "image_to_string", return_value="Office closes at 5 PM."
            ):
                self.assertEqual(
                    ocr.extract_text_from_image(str(image_path)),
                    "Office closes at 5 PM.",
                )

    @unittest.skipUnless(ocr.tesseract_available(), "Tesseract executable is unavailable")
    def test_real_tesseract_image_upload_is_indexed_and_retrievable(self):
        with tempfile.TemporaryDirectory() as directory:
            environment = os.environ.copy()
            environment.update({
                "PYTHONPATH": str(ROOT / "backend"),
                "DATA_DIR": str(Path(directory) / "data"),
                "CHROMA_PERSIST_DIR": str(Path(directory) / "chroma"),
                "UPLOAD_DIR": str(Path(directory) / "documents"),
                "EMBEDDING_PROVIDER": "local",
                "EMBEDDING_MODEL": "local-hash-v1",
                "LLM_PROVIDER": "local",
            })
            script = """
import asyncio
import io
import os
from PIL import Image, ImageDraw, ImageFont
from fastapi import UploadFile
from app.routes.documents import upload_documents
from app.services.document_manager import document_manager
from app.services.retriever import retriever

image = Image.new("RGB", (1600, 260), "white")
draw = ImageDraw.Draw(image)
font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
font = ImageFont.truetype(font_path, 70) if os.path.isfile(font_path) else ImageFont.load_default()
draw.text((40, 60), "Annual leave balance is 18 days", fill="black", font=font)
image_bytes = io.BytesIO()
image.save(image_bytes, format="PNG")
uploaded = asyncio.run(upload_documents([
    UploadFile(filename="leave-notice.png", file=io.BytesIO(image_bytes.getvalue()))
]))
assert uploaded[0].upload_status == "SUCCESS", uploaded[0].error_message
document = document_manager.get_document(uploaded[0].id)
assert document and document.chunks
extracted = " ".join(chunk.text for chunk in document.chunks).lower()
assert "annual leave balance is 18 days" in extracted, extracted
sources = retriever.retrieve("annual leave balance 18 days")
assert sources and sources[0].filename == "leave-notice.png"
assert "18 days" in sources[0].text.lower(), sources[0].text
"""
            completed = subprocess.run(
                [PYTHON, "-c", script],
                cwd=ROOT,
                env=environment,
                capture_output=True,
                text=True,
                timeout=120,
            )
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)

    def test_embedding_profile_mismatch_fails_without_deleting_index(self):
        with tempfile.TemporaryDirectory() as directory:
            store = VectorStore(
                persist_dir=str(Path(directory) / "chroma"),
                collection_name="profile_test",
            )
            original_count = store.collection.count()
            with patch.object(settings, "EMBEDDING_PROVIDER", "openai"), patch.object(
                settings, "EMBEDDING_MODEL", "text-embedding-3-small"
            ), patch.object(settings, "OPENAI_API_KEY", "test-key"):
                with self.assertRaisesRegex(RuntimeError, "active index"):
                    store.ensure_compatible()
            self.assertEqual(store.collection.count(), original_count)

    def test_external_embedding_failure_is_explicit_and_never_uses_local_vectors(self):
        with patch.object(settings, "EMBEDDING_PROVIDER", "openai"), patch.object(
            settings, "EMBEDDING_MODEL", "text-embedding-3-small"
        ), patch.object(settings, "OPENAI_API_KEY", "test-key"), patch.object(
            embedding_function,
            "_embed_openai",
            side_effect=RuntimeError("provider unavailable"),
        ), patch.object(
            embedding_function,
            "_embed_local",
        ) as local_embeddings:
            with self.assertRaisesRegex(EmbeddingConfigurationError, "OpenAI embedding request failed"):
                embedding_function(["annual paid leave"])
        local_embeddings.assert_not_called()


class ConflictRegressionTests(unittest.TestCase):
    def test_related_numeric_and_polarity_conflicts(self):
        detector = ConflictDetector()
        cases = [
            (
                "How many days are available for remote work?",
                "Remote work is allowed for 2 days each week.",
                "Remote work is allowed for 3 days each week.",
            ),
            (
                "Is remote work permitted?",
                "Remote work is permitted for employees.",
                "Remote work is prohibited for employees.",
            ),
            (
                "Is MFA required?",
                "MFA is required for all employees.",
                "MFA is not required for employees.",
            ),
            (
                "Is dental coverage included?",
                "Dental coverage is included.",
                "Dental coverage is excluded.",
            ),
            (
                "Does the company provide pet insurance?",
                "The company provides pet insurance.",
                "The company does not provide pet insurance.",
            ),
            (
                "Is a remote-work allowance available?",
                "Remote-work allowance: yes.",
                "Remote-work allowance: no.",
            ),
            (
                "How long is the probation period?",
                "The probation period lasts 3 months.",
                "The probation period lasts 6 months.",
            ),
        ]
        for index, (question, first, second) in enumerate(cases):
            with self.subTest(question=question):
                result = detector.check_conflicts(
                    question,
                    [
                        reference("Document A", first, f"a-{index}"),
                        reference("Document B", second, f"b-{index}"),
                    ],
                )
                self.assertTrue(result.conflict_detected)
                self.assertEqual(len(result.conflicting_sources), 2)

    def test_unrelated_numeric_values_do_not_create_a_conflict(self):
        result = ConflictDetector().check_conflicts(
            "How many remote work days are allowed?",
            [
                reference(
                    "Document A",
                    "Remote work is allowed for 2 days each week. Passwords rotate every 90 days.",
                    "a",
                ),
                reference(
                    "Document B",
                    "Remote work is allowed for 2 days each week. Passwords rotate every 180 days.",
                    "b",
                ),
            ],
        )
        self.assertFalse(result.conflict_detected)


class ProcessRestartTests(unittest.TestCase):
    def _environment(self, directory: str) -> dict[str, str]:
        environment = os.environ.copy()
        environment.update({
            "PYTHONPATH": str(ROOT / "backend"),
            "DATA_DIR": str(Path(directory) / "data"),
            "CHROMA_PERSIST_DIR": str(Path(directory) / "chroma"),
            "UPLOAD_DIR": str(Path(directory) / "documents"),
            "EMBEDDING_PROVIDER": "local",
            "EMBEDDING_MODEL": "local-hash-v1",
            "LLM_PROVIDER": "local",
        })
        return environment

    def _run_python(self, script: str, environment: dict[str, str]) -> None:
        completed = subprocess.run(
            [PYTHON, "-c", script],
            cwd=ROOT,
            env=environment,
            capture_output=True,
            text=True,
            timeout=120,
        )
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)

    def test_index_details_retrieval_and_history_survive_process_restart(self):
        with tempfile.TemporaryDirectory() as directory:
            environment = self._environment(directory)
            self._run_python(
                f"""
import asyncio
import io
import docx
import pymupdf
from fastapi import UploadFile
from PIL import Image
from app.services.document_manager import document_manager
from app.models.investigation import InvestigationRequest
from app.routes.documents import upload_documents
from app.routes.investigation import ask_question
image = io.BytesIO()
Image.new("RGB", (20, 20), "white").save(image, format="PNG")
pdf = pymupdf.open()
pdf.new_page().insert_text((72, 72), "A PDF policy describes a documented leave allowance.")
pdf_bytes = io.BytesIO()
pdf.save(pdf_bytes)
word = docx.Document()
word.add_paragraph("A DOCX handbook describes a documented leave allowance.")
docx_bytes = io.BytesIO()
word.save(docx_bytes)
uploads = asyncio.run(upload_documents([
    UploadFile(filename="stipend.txt", file=io.BytesIO(
        b"The annual library stipend is $700 for each employee."
    )),
    UploadFile(filename="policy.pdf", file=io.BytesIO(pdf_bytes.getvalue())),
    UploadFile(filename="handbook.docx", file=io.BytesIO(docx_bytes.getvalue())),
    UploadFile(filename="empty.txt", file=io.BytesIO(b"")),
    UploadFile(filename="legacy.doc", file=io.BytesIO(b"not a DOCX")),
    UploadFile(filename="blank.png", file=io.BytesIO(image.getvalue())),
]))
assert [item.upload_status for item in uploads] == [
    "SUCCESS", "SUCCESS", "SUCCESS", "FAILED", "FAILED", "FAILED"
]
assert uploads[0].chunks_count > 0
result = asyncio.run(ask_question(InvestigationRequest(
    question="What is the annual library stipend?"
)))
assert result.sources and result.sources[0].chunk_id
assert "$700" in result.answer
""",
                environment,
            )

            self._run_python(
                """
import asyncio
from app.services.document_manager import document_manager
from app.services.retriever import retriever
from app.models.investigation import InvestigationRequest
from app.routes.investigation import ask_question, get_investigation, share_investigation
doc_info = next(
    doc for doc in document_manager.get_all_documents()
    if doc.filename == "stipend.txt"
)
doc = document_manager.get_document(doc_info.id)
assert doc is not None and doc.chunks and "$700" in doc.chunks[0].text
retrieved = retriever.retrieve("annual library stipend")
assert retrieved and retrieved[0].filename == "stipend.txt"
answer = asyncio.run(ask_question(InvestigationRequest(
    question="What is the annual library stipend?"
)))
assert "$700" in answer.answer and answer.sources
history = document_manager.get_investigations()
assert len(history) == 2 and history[0].sources and history[1].sources
opened = asyncio.run(get_investigation(history[1].id))
assert opened.sources[0].chunk_id == history[1].sources[0].chunk_id
share = asyncio.run(share_investigation(history[1].id))
assert share.share_url == f"/#investigation-{history[1].id}"
""",
                environment,
            )

    def test_staged_leave_documents_transition_from_supported_answer_to_conflict(self):
        with tempfile.TemporaryDirectory() as directory:
            environment = self._environment(directory)
            self._run_python(
                """
import asyncio
import io
import pymupdf
from fastapi import UploadFile
from app.models.investigation import InvestigationRequest
from app.routes.documents import upload_documents
from app.routes.investigation import ask_question

def pdf_bytes(text):
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    content = io.BytesIO()
    document.save(content)
    document.close()
    return content.getvalue()

question = "How many paid leaves can an employee take per year?"
initial = asyncio.run(upload_documents([
    UploadFile(filename="Employee_Handbook.pdf", file=io.BytesIO(
        pdf_bytes("Annual paid leave allowance is 18 days.")
    )),
    UploadFile(filename="HR_Rules.txt", file=io.BytesIO(
        b"HR requests must be sent to the employee's department."
    )),
]))
assert [item.upload_status for item in initial] == ["SUCCESS", "SUCCESS"]
first_answer = asyncio.run(ask_question(InvestigationRequest(question=question)))
assert not first_answer.conflict_detected
assert "18 days" in first_answer.answer
assert first_answer.sources[0].filename == "Employee_Handbook.pdf"

conflict_upload = asyncio.run(upload_documents([
    UploadFile(filename="Leave_Policy.pdf", file=io.BytesIO(
        pdf_bytes("Annual paid leave allowance is 20 days.")
    )),
]))
assert conflict_upload[0].upload_status == "SUCCESS"
conflict_answer = asyncio.run(ask_question(InvestigationRequest(question=question)))
assert conflict_answer.conflict_detected
assert conflict_answer.confidence == "LOW"
assert {source.filename for source in conflict_answer.sources} == {
    "Employee_Handbook.pdf", "Leave_Policy.pdf"
}
assert "18 days" in conflict_answer.answer and "20 days" in conflict_answer.answer
""",
                environment,
            )

    def test_seed_script_is_idempotent_and_reset_preserves_user_documents(self):
        with tempfile.TemporaryDirectory() as directory:
            environment = self._environment(directory)
            seed = [PYTHON, str(ROOT / "seed_demo.py")]
            for _ in range(2):
                completed = subprocess.run(
                    seed, cwd=ROOT, env=environment, capture_output=True, text=True, timeout=120
                )
                self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)

            registry_path = Path(directory) / "data" / "documents.json"
            registry = json.loads(registry_path.read_text(encoding="utf-8"))
            names = [item["filename"] for item in registry]
            self.assertEqual(len(names), len(set(names)))

            user_document = dict(registry[0])
            user_document.update({
                "id": "user-document",
                "filename": "user-notes.txt",
                "status": "failed",
                "upload_status": "FAILED",
                "chunks_count": 0,
            })
            registry.append(user_document)
            registry_path.write_text(json.dumps(registry), encoding="utf-8")

            completed = subprocess.run(
                seed + ["--reset"],
                cwd=ROOT,
                env=environment,
                capture_output=True,
                text=True,
                timeout=120,
            )
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            registry = json.loads(registry_path.read_text(encoding="utf-8"))
            names = [item["filename"] for item in registry]
            self.assertEqual(len(names), len(set(names)))
            self.assertIn("user-notes.txt", names)


class ConfigurationSafetyTests(unittest.TestCase):
    def test_settings_never_return_keys_or_crash_on_invalid_embedding_profile(self):
        with patch.object(settings, "EMBEDDING_PROVIDER", "invalid-provider"):
            response = asyncio.run(system_routes.get_settings())
        self.assertFalse(response["vector_store_ready"])
        self.assertIsNone(response["embedding_dimension"])
        self.assertNotIn("openai_api_key", response)
        self.assertNotIn("gemini_api_key", response)

    def test_settings_api_rejects_browser_submitted_api_keys(self):
        with self.assertRaises(Exception):
            system_routes.SettingsUpdateRequest.model_validate({
                "llm_provider": "local",
                "openai_api_key": "must-not-be-accepted",
            })

    def test_destructive_reset_is_disabled_in_production(self):
        with patch.object(settings, "APP_ENV", "production"), patch.object(
            settings, "ALLOW_DESTRUCTIVE_RESET", True
        ), patch.object(system_routes.vector_store, "clear_all") as clear_all:
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(system_routes.reset_database())
        self.assertEqual(raised.exception.status_code, 403)
        clear_all.assert_not_called()

    def test_production_runtime_settings_updates_are_disabled(self):
        with patch.object(settings, "APP_ENV", "production"):
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(system_routes.update_settings(
                    system_routes.SettingsUpdateRequest(
                        llm_provider="local",
                        chunk_size=500,
                        chunk_overlap=100,
                    )
                ))
        self.assertEqual(raised.exception.status_code, 403)


class ProviderFailureTests(unittest.TestCase):
    def test_selected_provider_failure_does_not_return_local_fallback(self):
        generator = AnswerGenerator()
        with patch.object(settings, "LLM_PROVIDER", "openai"), patch.object(
            settings, "OPENAI_API_KEY", "test-key"
        ), patch.object(
            generator,
            "_generate_openai_answer",
            side_effect=RuntimeError("provider unavailable"),
        ), patch.object(
            generator,
            "_generate_local_grounded_answer",
        ) as local_fallback:
            with self.assertRaisesRegex(AnswerProviderError, "no local fallback"):
                generator.generate_investigation(
                    "provider-test",
                    "What is the annual leave policy?",
                    [reference("policy.txt", "Annual leave policy grants 20 days.", "leave")],
                    "now",
                )
        local_fallback.assert_not_called()

    def test_external_answer_provider_requires_backend_key(self):
        with patch.object(settings, "LLM_PROVIDER", "gemini"), patch.object(
            settings, "GEMINI_API_KEY", ""
        ):
            with self.assertRaisesRegex(AnswerProviderError, "GEMINI_API_KEY"):
                AnswerGenerator().generate_investigation(
                    "provider-test",
                    "What is the annual leave policy?",
                    [reference("policy.txt", "Annual leave policy grants 20 days.", "leave")],
                    "now",
                )

    def test_gemini_api_key_is_sent_in_header_not_query_string(self):
        class Response:
            def raise_for_status(self):
                return None

            def json(self):
                return {"candidates": [{"content": {"parts": [{"text": "20 days."}]}}]}

        class Client:
            def __init__(self, **kwargs):
                pass

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def post(self, url, **kwargs):
                self.url = url
                self.headers = kwargs["headers"]
                return Response()

        client = Client()
        with patch("app.services.answer_generator.httpx.Client", return_value=client), patch.object(
            settings, "GEMINI_API_KEY", "test-secret"
        ):
            AnswerGenerator()._generate_gemini_answer(
                "provider-test",
                "How many paid leave days?",
                [reference("leave.txt", "Annual paid leave allowance is 20 days.", "leave")],
                "now",
            )
        self.assertNotIn("test-secret", client.url)
        self.assertEqual(client.headers["x-goog-api-key"], "test-secret")


if __name__ == "__main__":
    unittest.main()
